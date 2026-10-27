"""Internal technology-lookup tool: paste domains, get a technology report."""
import io
import hmac
import multiprocessing
import os
import threading
import time
import uuid
import webbrowser
from concurrent.futures import ProcessPoolExecutor, as_completed

from flask import Flask, jsonify, request, send_file, send_from_directory
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import store
from analysis import aggregate, analyze
from detector import Detector
from scanner import init_worker, parse_domains, scan_in_worker

app = Flask(__name__, static_folder="static")
APP_PASSWORD = os.environ.get("APP_PASSWORD", "")
APP_USERNAME = os.environ.get("APP_USERNAME", "buildwith")


@app.before_request
def require_authentication():
    """Protect the public deployment; local development stays password-free."""
    if not APP_PASSWORD:
        if os.environ.get("RAILWAY_SERVICE_NAME") or os.environ.get("RENDER_SERVICE_ID"):
            return jsonify(error="APP_PASSWORD must be configured before serving requests"), 503
        return None
    auth = request.authorization
    if (auth and hmac.compare_digest(auth.username or "", APP_USERNAME)
            and hmac.compare_digest(auth.password or "", APP_PASSWORD)):
        return None
    return jsonify(error="authentication required"), 401, {"WWW-Authenticate": 'Basic realm="BuildWith"'}


detector = None  # main-process copy, only used for /api/info
# Worker processes re-import this module under multiprocessing's spawn; skip one-time setup there.
if multiprocessing.current_process().name == "MainProcess":
    store.init()
    detector = Detector()
pool_lock = threading.Lock()
POOL = None
JOBS = {}  # scan_id -> live scan dict (while running)
WORKERS = int(os.environ.get("WORKERS", "6"))  # CPU-bound regex matching: processes, not threads (GIL)


def _ts(epoch):
    return time.strftime("%d/%m/%Y %H:%M", time.localtime(epoch)) if epoch else ""


def get_pool():
    global POOL
    with pool_lock:
        if POOL is None:
            POOL = ProcessPoolExecutor(WORKERS, initializer=init_worker)
        return POOL


def _diff(result, scan_id):
    prev, prev_at = store.previous_result(result["domain"], scan_id)
    if not prev or not result["ok"]:
        return None
    old = {t["name"] for t in prev["technologies"]}
    new = {t["name"] for t in result["technologies"]}
    return {"added": sorted(new - old), "removed": sorted(old - new), "prev_at": prev_at}


def _run_scan(scan_id, domains):
    live = JOBS[scan_id]
    pool = get_pool()
    futures = {pool.submit(scan_in_worker, d): d for d in domains}
    for f in as_completed(futures):
        d = futures[f]
        try:
            res = f.result()
        except Exception as e:
            res = {"domain": d, "ok": False, "error": f"{type(e).__name__}: {e}",
                   "technologies": [], "dns": {}, "containers": [], "title": "",
                   "status_code": None, "final_url": None}
        res["changes"] = _diff(res, scan_id)
        store.add_result(scan_id, res)
        live["results"].append(res)
        live["done"] += 1
    live["status"] = "done"
    store.finish_scan(scan_id)


@app.get("/")
def index():
    return send_from_directory("static", "index.html")


@app.post("/api/scans")
def create_scan():
    domains = parse_domains((request.get_json(silent=True) or {}).get("domains", ""))
    if not domains:
        return jsonify(error="Không tìm thấy domain hợp lệ"), 400
    scan_id = uuid.uuid4().hex[:10]
    store.create_scan(scan_id, len(domains))
    JOBS[scan_id] = {"id": scan_id, "created_at": time.time(), "total": len(domains),
                     "done": 0, "status": "running", "results": []}
    threading.Thread(target=_run_scan, args=(scan_id, domains), daemon=True).start()
    return jsonify(id=scan_id, total=len(domains))


@app.get("/api/scans")
def list_scans():
    return jsonify(store.list_scans())


def _load(scan_id):
    """Scan with per-domain analysis and the cross-domain aggregate attached."""
    scan = JOBS.get(scan_id) or store.get_scan(scan_id)
    if not scan:
        return None
    results = [dict(r, analysis=analyze(r)) for r in list(scan["results"])]
    return dict(scan, results=results, aggregate=aggregate(results))


@app.get("/api/scans/<scan_id>")
def get_scan(scan_id):
    scan = _load(scan_id)
    if not scan:
        return jsonify(error="not found"), 404
    return jsonify(scan)


@app.get("/api/info")
def info():
    return jsonify(detector.sources | {"technologies": len(detector.techs)})


@app.post("/api/signatures/update")
def update_signatures():
    import update_signatures as us
    try:
        us.update()
    except Exception as e:
        return jsonify(error=f"{type(e).__name__}: {e}"), 502
    global POOL
    detector.load()
    with pool_lock:  # new scans use fresh workers; running scans finish on the old pool
        old, POOL = POOL, None
    if old:
        old.shutdown(wait=False)
    return jsonify(detector.sources | {"technologies": len(detector.techs)})


HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
WRAP = Alignment(wrap_text=True, vertical="top")
SIGNAL_TAG = {"onsite": "", "tagmanager": " [qua TM]", "linked": " [liên kết]", "dns": " [DNS]",
              "verification": " [xác thực domain]", "implied": " [suy luận]"}
QUALITY_FILL = {"full": "E5F6EC", "thin": "FDF0D5", "blocked": "FDE4E1", "http_error": "FDE4E1",
                "unreachable": "EEEEEE"}


def _sheet(wb, title, header, widths, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    return ws


def _finish(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    ws.auto_filter.ref = ws.dimensions


def _group_cell(a, group):
    return "\n".join(t["name"] + SIGNAL_TAG[t["signal"]] for t in a["techs"] if t["group"] == group)


@app.get("/api/scans/<scan_id>/export")
def export(scan_id):
    scan = _load(scan_id)
    if not scan:
        return jsonify(error="not found"), 404
    results = sorted(scan["results"], key=lambda r: r["domain"])
    agg = scan["aggregate"]
    wb = Workbook()

    # 1. How to read
    ws = _sheet(wb, "Cách đọc", ["Mục", "Giá trị", "Ý nghĩa"], [22, 26, 110], first=True)
    for k, s in agg["signals"].items():
        ws.append(["Loại tín hiệu", s["label"] + SIGNAL_TAG[k], s["help"]])
    for lvl, help_ in agg["levels"].items():
        ws.append(["Độ tin cậy", lvl, help_])
    ws.append(["Tình trạng quét", "Đầy đủ / Ít dữ liệu / Bị chặn / Không truy cập được",
               "Cho biết kết quả của domain đó đáng tin tới đâu. Domain “Bị chặn” hoặc “Không truy cập được” "
               "thường thiếu nhiều công nghệ: không có nghĩa là họ không dùng."])
    ws.append(["Phạm vi", f"Quét lúc {_ts(scan.get('created_at'))}",
               "Mỗi domain: trang chủ + tối đa 5 trang con (liên hệ, báo giá, đại lý, demo, newsletter...), "
               "container tag manager (GTM, Adobe Launch, Tealium) và DNS (MX, TXT, NS, CNAME). "
               "Không chạy JavaScript, không đăng nhập."])
    _finish(ws)

    # 2. Overview, one row per domain
    ws = _sheet(wb, "Tổng quan",
                ["Domain", "Tình trạng quét", "Tóm tắt", "Nền tảng website", "Martech & dữ liệu KH",
                 "CRM, bán hàng & CSKH", "Email, DNS & tài khoản", "Quảng cáo", "Salesforce Org ID",
                 "Thay đổi so với lần trước", "Title", "Trang đã quét"],
                [26, 18, 70, 30, 34, 34, 30, 26, 26, 26, 34, 50])
    for r in results:
        a = r["analysis"]
        ch = r.get("changes")
        change = ("lần đầu quét" if not ch else
                  "\n".join([f"+ {n}" for n in ch["added"]] + [f"− {n}" for n in ch["removed"]]) or "không đổi")
        orgs = "\n".join(f"{o['id']}: {'; '.join(o['sources'])}" for o in r.get("salesforce_org_ids") or [])
        ws.append([r["domain"], f"{a['quality']['label']}\n{a['quality']['note']}", "\n".join(f"• {s}" for s in a["summary"]),
                   _group_cell(a, "platform"), _group_cell(a, "martech"), _group_cell(a, "crm"),
                   _group_cell(a, "email"), _group_cell(a, "ads"), orgs, change, r.get("title") or "",
                   "\n".join(r.get("pages") or [])])
        ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=QUALITY_FILL.get(a["quality"]["key"], "FFFFFF"))
    _finish(ws)

    # 3. Every detection with its evidence
    ws = _sheet(wb, "Chi tiết công nghệ",
                ["Domain", "Nhóm", "Công nghệ", "Hãng", "Loại tín hiệu", "Độ tin cậy", "Ở đâu", "Bằng chứng",
                 "Danh mục"], [26, 26, 36, 14, 20, 12, 30, 80, 30])
    for r in results:
        for t in r["analysis"]["techs"]:
            where = ", ".join(t["where"]) or ", ".join(t["via"])
            ev = "\n".join(f"{e['label']}{' [' + e['page'] + ']' if e['page'] else ''}"
                           f"{' (' + e['via'] + ')' if e.get('via') else ''}: {e['value']}" for e in t["evidence"])
            ws.append([r["domain"], t["group_label"], t["name"], t["vendor"] or "", t["signal_label"], t["level"],
                       where, ev, ", ".join(t["categories"])])
    _finish(ws)

    # 4. Vendors across domains
    ws = _sheet(wb, "Theo hãng", ["Hãng", "Số domain", "Có trên website", "Chỉ DNS / liên kết / xác thực",
                                  "Sản phẩm (số domain)", "Domain có trên website", "Domain chỉ DNS / liên kết"],
                [16, 10, 14, 18, 60, 50, 50])
    for v in agg["vendors"]:
        on = sorted(d for d, s in v["domains"].items() if s in ("onsite", "tagmanager"))
        off = sorted(d for d in v["domains"] if d not in on)
        prods = "\n".join(f"{p} ({len(ds)})" for p, ds in sorted(v["products"].items(), key=lambda kv: -len(kv[1])))
        ws.append([v["vendor"], len(v["domains"]), len(on), len(off), prods, "\n".join(on), "\n".join(off)])
    _finish(ws)

    # 5. Technologies across domains
    ws = _sheet(wb, "Theo công nghệ", ["Công nghệ", "Nhóm", "Hãng", "Số domain", "Trên website", "Qua tag manager",
                                       "Tín hiệu khác", "Domains"], [36, 26, 14, 10, 12, 14, 14, 70])
    for t in agg["techs"]:
        sig = list(t["domains"].values())
        ws.append([t["name"], t["group_label"], t["vendor"] or "", len(sig), sig.count("onsite"),
                   sig.count("tagmanager"), len(sig) - sig.count("onsite") - sig.count("tagmanager"),
                   "\n".join(f"{d}{SIGNAL_TAG[s]}" for d, s in sorted(t["domains"].items()))])
    _finish(ws)

    # 6. Raw DNS
    ws = _sheet(wb, "DNS", ["Domain", "Loại", "Giá trị"], [26, 8, 110])
    for r in results:
        for rtype, vals in (r.get("dns") or {}).items():
            for val in vals:
                ws.append([r["domain"], rtype, val])
    _finish(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"tech-report-{scan_id}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    threading.Timer(1.0, lambda: webbrowser.open("http://127.0.0.1:5000")).start()
    app.run(host="127.0.0.1", port=5000, threaded=True)
